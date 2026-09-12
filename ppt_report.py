"""Generate one chart-pair slide per qualifying defect/process combination."""
from pathlib import Path

from matplotlib.figure import Figure
from matplotlib.backends.backend_agg import FigureCanvasAgg

from defect_worse_tool import (
    read_table, add_grouping_columns, build_worse_tool_result,
    filter_by_recent_scan_time, handle_outliers_for_defect,
    apply_special_process_rules, apply_process_aggregation,
)


class Value:
    """Minimal variable adapter for headless chart rendering; never touches Tk."""
    def __init__(self, value):
        self.value = value

    def get(self):
        return self.value

    def set(self, value):
        self.value = value


def make_renderer():
    from defect_worse_ui import DefectWorseToolApp
    renderer = object.__new__(DefectWorseToolApp)
    renderer.fig = Figure(figsize=(10, 7), dpi=140)
    renderer.canvas = FigureCanvasAgg(renderer.fig)
    defaults = dict(box_line_width=1.4, box_label_font_size=0.0,
                    show_box_count=True, show_box_median=True, show_box_mean=True,
                    line_width=1.6, marker_size=3.0, color_scheme="Distinct",
                    custom_color="#1565C0", y_min="", y_max="",
                    selected_chart_item="", status="")
    defaults.update(show_bsl_line=True, show_threshold_line=True, show_golden_line=True)
    for name, value in defaults.items():
        setattr(renderer, name, Value(value))
    renderer.chart_artist_registry = {}
    renderer.artist_style_overrides = {}
    renderer.selected_chart_artist = None
    return renderer


def generate_report(context, log_callback=None):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from defect_worse_ui import prepare_trend_data
    from chart_context import annotate_chart_data, analysis_references
    from PIL import Image

    log = log_callback or (lambda message: None)
    log("Recalculating worse tools using current analysis settings.")
    result = build_worse_tool_result(
        context.raw_data_path, context.bsl_path or None,
        input_sheet=context.input_sheet, defect_columns=context.defect_columns,
        bsl_source=context.bsl_source, data_window=context.data_window,
        bsl_multiplier=context.bsl_multiplier, min_wafers=context.min_wafers,
        outlier_sigma=context.outlier_sigma, outlier_handling=context.outlier_handling,
        process_aggregation=context.process_aggregation,
        special_process_rules=context.special_process_rules,
    )
    if result.empty:
        raise ValueError("No worse tools match the current settings; no PPT was written.")
    raw = add_grouping_columns(read_table(context.raw_data_path, context.input_sheet))
    output = Path(context.ppt_output_path)
    if context.ppt_template_path and output.resolve() == Path(context.ppt_template_path).resolve():
        raise ValueError("PPT output must differ from the template path.")
    images = Path(context.input_image_path) if context.input_image_path else output.parent / (output.stem + "_charts")
    images.mkdir(parents=True, exist_ok=True)
    presentation = Presentation(context.ppt_template_path) if context.ppt_template_path else Presentation()
    if not context.ppt_template_path:
        presentation.slide_width = Inches(16)
        presentation.slide_height = Inches(9)
    # Retain template masters/themes and replace its example slides in memory.
    for slide_id in list(presentation.slides._sldIdLst):
        presentation.part.drop_rel(slide_id.rId)
        presentation.slides._sldIdLst.remove(slide_id)
    layout = min(presentation.slide_layouts, key=lambda candidate: len(candidate.placeholders))
    renderer = make_renderer()
    for name in ("show_bsl_line", "show_threshold_line", "show_golden_line"):
        renderer.__dict__[name].set(getattr(context, name, True))
    pages = result[["Defect type", "Stage_ID", "Step_ID"]].drop_duplicates()
    width, height = presentation.slide_width, presentation.slide_height
    for page_number, (_, row) in enumerate(pages.iterrows(), 1):
        defect, stage, step = str(row["Defect type"]), str(row["Stage_ID"]), str(row["Step_ID"])
        layer = step if stage in {"ALL_STAGES", "SPECIAL_STEP_ONLY"} else "{} / {}".format(stage, step)
        paths = []
        options = dict(vars(context))
        references = analysis_references(raw, defect, options)
        for chart_kind, window in (("box", "14d"), ("trend", "all")):
            data = filter_by_recent_scan_time(raw, window)
            data = handle_outliers_for_defect(data, defect, context.outlier_sigma, context.outlier_handling)
            data = apply_special_process_rules(data, defect, context.special_process_rules)
            data = apply_process_aggregation(data, context.process_aggregation)
            data = data.loc[(data["Stage_ID"].astype(str) == stage) & (data["Step_ID"].astype(str) == step)]
            data = renderer._filter_chart_group_mode(data, context.chart_group_mode)
            if not data.empty:
                annotate_chart_data(renderer, raw, data, defect, options, window, references)
            if data.empty:
                renderer.fig.clear()
                ax = renderer.fig.add_subplot(111)
                ax.set_axis_off()
                ax.text(0.5, 0.5, "No data in latest 14 days" if chart_kind == "box" else "No data",
                        ha="center", va="center", transform=ax.transAxes)
            elif chart_kind == "box":
                renderer._draw_box(defect, layer, data)
            else:
                trend = prepare_trend_data(data, defect, context.time_column)
                renderer._draw_trend_sequence_by_tool(defect, layer, context.time_column, trend)
            path = images / ("{}_page_{:03d}_{}.png".format(output.stem, page_number, chart_kind))
            if not data.empty:
                renderer.export_current_chart(str(path))
            else:
                renderer.fig.savefig(str(path), dpi=140)
            paths.append(path)
        slide = presentation.slides.add_slide(layout)
        for shape in list(slide.placeholders):
            shape._element.getparent().remove(shape._element)
        title = slide.shapes.add_textbox(int(width * .04), int(height * .025), int(width * .92), int(height * .10))
        paragraph = title.text_frame.paragraphs[0]
        paragraph.text = "{} cross to {}".format(defect, layer)
        paragraph.font.size = Pt(24)
        for index, path in enumerate(paths):
            left = int(width * (.02 + .5 * index))
            picture_width = int(width * .46)
            with Image.open(path) as img:
                aspect = img.height / img.width
            picture_height = int(picture_width * aspect)
            available_height = int(height * .73)
            if picture_height > available_height:
                picture_width = int(available_height / aspect)
            slide.shapes.add_picture(str(path), left, int(height * .18), width=picture_width)
            caption = slide.shapes.add_textbox(left, int(height * .13), int(width * .46), int(height * .05))
            caption.text = "Box | latest 14 days" if index == 0 else "Sequential trend | all data"
        log("Slide {}/{}: {} cross to {}".format(page_number, len(pages), defect, layer))
    output.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(str(output))
    log("Saved {} slides: {}".format(len(pages), output))
    return output
