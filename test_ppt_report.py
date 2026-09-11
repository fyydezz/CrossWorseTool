import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import pandas as pd
from pptx import Presentation

from defect_worse_tool import build_worse_tool_result
from defect_worse_ui import rank_worse_results
from ppt_integration import PPTGenerationContext, run_ppt_generation
from ppt_report import make_renderer


class ReportTests(unittest.TestCase):
    def test_recent_baseline_is_independent_of_analysis_window(self):
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / "raw.csv"
            rows = []
            for tool, count, time in (("KP01", 100, "2026-06-01"),
                                      ("KP02", 10, "2026-08-20"),
                                      ("KP03", 0, "2026-08-20")):
                for wafer in range(5):
                    rows.append(dict(Lot_ID=tool, Wafer_NO=wafer, Scan_Time=time,
                                     Stage_ID="S", Step_ID="P", Equipment_ID=tool,
                                     Chamber_ID="C", D1=count))
            pd.DataFrame(rows).to_csv(source, index=False)
            result = build_worse_tool_result(str(source), bsl_source="recent_mean", data_window="all")
            self.assertEqual(set(result["BSL count"]), {5.0})
            self.assertIn("KP01", result["Equipment ID"].tolist())

    def test_priority_favors_severity_and_sample_size(self):
        data = pd.DataFrame({"Mean_Count": [10, 20, 20], "BSL count": [5, 5, 5],
                             "Wafer_Count": [5, 5, 20]})
        self.assertEqual(rank_worse_results(data).index.tolist(), [2, 1, 0])

    def test_sparse_box_keeps_inline_statistics(self):
        renderer = make_renderer()
        data = pd.DataFrame({"Chart_Group": ["A"] * 6 + ["B"] * 6,
                             "D1": list(range(6)) * 2})
        renderer._draw_box("D1", "S", data)
        self.assertEqual(len(renderer.fig.axes[0].texts), 2)
        self.assertIn("Mean=", renderer.fig.axes[0].texts[0].get_text())

    def test_report_deduplicates_layer_and_uses_separate_windows(self):
        with tempfile.TemporaryDirectory() as folder:
            output = Path(folder) / "report.pptx"
            context = PPTGenerationContext(
                str(output), "", str(Path(folder) / "images"),
                "demo_defect_data.csv", "demo_bsl.csv", "", None, "worse_tool",
                None, 1.5, 5, 3.0, None, None,
            )
            from ppt_report import filter_by_recent_scan_time
            windows = []
            def capture(data, window):
                windows.append(window)
                return filter_by_recent_scan_time(data, window)
            with patch("ppt_report.filter_by_recent_scan_time", side_effect=capture):
                run_ppt_generation(context)
            deck = Presentation(str(output))
            self.assertGreater(len(deck.slides), 0)
            titles = []
            for slide in deck.slides:
                self.assertEqual(sum(shape.shape_type == 13 for shape in slide.shapes), 2)
                titles.extend(shape.text for shape in slide.shapes
                              if shape.has_text_frame and " cross to " in shape.text)
            self.assertEqual(len(titles), len(set(titles)))
            self.assertEqual(windows, ["14d", "all"] * len(deck.slides))


if __name__ == "__main__":
    unittest.main()
